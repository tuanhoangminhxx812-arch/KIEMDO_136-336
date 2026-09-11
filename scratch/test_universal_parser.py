import openpyxl
import xlrd
import io
import re
import os
import sys

def parse_num(v):
    if v is None or v == '': return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(' ', '').replace('\xa0', '').strip()
    if s.startswith('(') and s.endswith(')'):
        return -float(s[1:-1])
    try:
        return float(s)
    except:
        return 0.0

def parse_account_file_universal(fpath_or_stream):
    if not fpath_or_stream:
        return {}, []
    if isinstance(fpath_or_stream, str) and not os.path.exists(fpath_or_stream):
        return {}, []
        
    is_stream = not isinstance(fpath_or_stream, str)
    stream_bytes = None
    if is_stream:
        stream_bytes = fpath_or_stream.read()
        fpath_or_stream.seek(0)
        
    # Check if xls
    is_xls = False
    if is_stream:
        if stream_bytes and stream_bytes[:4] == b'\xd0\xcf\x11\xe0':
            is_xls = True
    elif isinstance(fpath_or_stream, str) and fpath_or_stream.lower().endswith('.xls'):
        is_xls = True
        
    accounts = {}
    transactions = []
    current_acc = None
    
    if not is_xls:
        try:
            target = io.BytesIO(stream_bytes) if is_stream else fpath_or_stream
            wb = openpyxl.load_workbook(target, data_only=True)
            ws = wb.active
            for r in range(1, ws.max_row + 1):
                c1 = ws.cell(r, 1).value
                if c1 and str(c1).startswith("Tài khoản:"):
                    acc_str = str(c1).replace("Tài khoản:", "").strip()
                    code_match = re.match(r'^(\d+)', acc_str)
                    if code_match:
                        current_acc = code_match.group(1)
                        accounts[current_acc] = {
                            'code': current_acc, 'title': acc_str,
                            'open_no': 0.0, 'open_co': 0.0,
                            'ps_no': 0.0, 'ps_co': 0.0,
                            'close_no': 0.0, 'close_co': 0.0,
                            'sub_records': []
                        }
                elif current_acc:
                    c1_str = str(c1) if c1 is not None else ""
                    if "Số dư đầu kỳ" in c1_str:
                        accounts[current_acc]['open_no'] = parse_num(ws.cell(r, 6).value)
                        accounts[current_acc]['open_co'] = parse_num(ws.cell(r, 7).value)
                    elif "Cộng phát sinh" in c1_str:
                        accounts[current_acc]['ps_no'] = parse_num(ws.cell(r, 6).value)
                        accounts[current_acc]['ps_co'] = parse_num(ws.cell(r, 7).value)
                    elif "Số dư cuối kỳ" in c1_str:
                        accounts[current_acc]['close_no'] = parse_num(ws.cell(r, 6).value)
                        accounts[current_acc]['close_co'] = parse_num(ws.cell(r, 7).value)
                    elif c1 and not c1_str.startswith("Khối:") and not c1_str.startswith("TỔNG") and not c1_str.startswith("CHI NHÁNH") and not c1_str.startswith("Nguồn bút toán") and not c1_str.startswith("SỔ CHI TIẾT") and not c1_str.startswith("Từ ngày:"):
                        tx_date = ws.cell(r, 2).value
                        tx_sub_doc = ws.cell(r, 3).value
                        tx_gl_doc = ws.cell(r, 4).value
                        tx_desc = ws.cell(r, 5).value
                        tx_no = parse_num(ws.cell(r, 6).value)
                        tx_co = parse_num(ws.cell(r, 7).value)
                        tx_creator = ws.cell(r, 9).value if ws.max_column >= 9 else ws.cell(r, 8).value
                        
                        tx_rec = {
                            'account': current_acc,
                            'source': c1,
                            'date': str(tx_date) if tx_date else '',
                            'sub_doc': str(tx_sub_doc) if tx_sub_doc else '',
                            'gl_doc': str(tx_gl_doc) if tx_gl_doc else '',
                            'desc': str(tx_desc) if tx_desc else '',
                            'ps_no': tx_no,
                            'ps_co': tx_co,
                            'creator': str(tx_creator) if tx_creator else ''
                        }
                        transactions.append(tx_rec)
                        accounts[current_acc]['sub_records'].append(tx_rec)
            return accounts, transactions
        except Exception as e:
            # Fallback to xlrd if openpyxl failed (could be misnamed .xls)
            is_xls = True

    if is_xls:
        try:
            if is_stream:
                wb = xlrd.open_workbook(file_contents=stream_bytes)
            else:
                wb = xlrd.open_workbook(fpath_or_stream)
            ws = wb.sheet_by_index(0)
            for r in range(ws.nrows):
                c1 = ws.cell_value(r, 0)
                if c1 and str(c1).startswith("Tài khoản:"):
                    acc_str = str(c1).replace("Tài khoản:", "").strip()
                    code_match = re.match(r'^(\d+)', acc_str)
                    if code_match:
                        current_acc = code_match.group(1)
                        accounts[current_acc] = {
                            'code': current_acc, 'title': acc_str,
                            'open_no': 0.0, 'open_co': 0.0,
                            'ps_no': 0.0, 'ps_co': 0.0,
                            'close_no': 0.0, 'close_co': 0.0,
                            'sub_records': []
                        }
                elif current_acc:
                    c1_str = str(c1) if c1 is not None else ""
                    if "Số dư đầu kỳ" in c1_str:
                        accounts[current_acc]['open_no'] = parse_num(ws.cell_value(r, 5) if ws.ncols > 5 else 0)
                        accounts[current_acc]['open_co'] = parse_num(ws.cell_value(r, 6) if ws.ncols > 6 else 0)
                    elif "Cộng phát sinh" in c1_str:
                        accounts[current_acc]['ps_no'] = parse_num(ws.cell_value(r, 5) if ws.ncols > 5 else 0)
                        accounts[current_acc]['ps_co'] = parse_num(ws.cell_value(r, 6) if ws.ncols > 6 else 0)
                    elif "Số dư cuối kỳ" in c1_str:
                        accounts[current_acc]['close_no'] = parse_num(ws.cell_value(r, 5) if ws.ncols > 5 else 0)
                        accounts[current_acc]['close_co'] = parse_num(ws.cell_value(r, 6) if ws.ncols > 6 else 0)
                    elif c1 and not c1_str.startswith("Khối:") and not c1_str.startswith("TỔNG") and not c1_str.startswith("CHI NHÁNH") and not c1_str.startswith("Nguồn bút toán") and not c1_str.startswith("SỔ CHI TIẾT") and not c1_str.startswith("Từ ngày:"):
                        tx_date = ws.cell_value(r, 1) if ws.ncols > 1 else ''
                        tx_sub_doc = ws.cell_value(r, 2) if ws.ncols > 2 else ''
                        tx_gl_doc = ws.cell_value(r, 3) if ws.ncols > 3 else ''
                        tx_desc = ws.cell_value(r, 4) if ws.ncols > 4 else ''
                        tx_no = parse_num(ws.cell_value(r, 5) if ws.ncols > 5 else 0)
                        tx_co = parse_num(ws.cell_value(r, 6) if ws.ncols > 6 else 0)
                        tx_creator = ws.cell_value(r, 8) if ws.ncols >= 9 else (ws.cell_value(r, 7) if ws.ncols >= 8 else '')
                        
                        tx_rec = {
                            'account': current_acc,
                            'source': c1,
                            'date': str(tx_date) if tx_date else '',
                            'sub_doc': str(tx_sub_doc) if tx_sub_doc else '',
                            'gl_doc': str(tx_gl_doc) if tx_gl_doc else '',
                            'desc': str(tx_desc) if tx_desc else '',
                            'ps_no': tx_no,
                            'ps_co': tx_co,
                            'creator': str(tx_creator) if tx_creator else ''
                        }
                        transactions.append(tx_rec)
                        accounts[current_acc]['sub_records'].append(tx_rec)
            return accounts, transactions
        except Exception as e:
            return {}, []
            
    return accounts, transactions

# Test universal parser on one xlsx
accs, txs = parse_account_file_universal('đầu vào/tháng 7/TK 136 - HCM.xlsx')
print(f"Universal parser test: {len(accs)} accounts, {len(txs)} transactions")
